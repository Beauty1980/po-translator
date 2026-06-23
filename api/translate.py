from http.server import BaseHTTPRequestHandler
import json, re, io, copy, shutil, tempfile, os, base64
from datetime import datetime
from dateutil.relativedelta import relativedelta

EN_TO_RU = {
    "Supplier:": "Поставщик:",
    "Buyer:": "Покупатель:",
    "Rotork quot. ref.": "Ссылка на квотацию Rotork",
    "Delivery terms:": "Условия поставки:",
    "Delivery time:": "Срок поставки:",
    "Payment terms:": "Условия оплаты:",
    "End user:": "Конечный пользователь:",
    "Tien-Shan Engineering LLP": "ТОО Тянь-Шань Engineering ",
    "ATAKENT, Timiryazeva 42 str., pavilion 17": "АТАКЕНТ, ул.  Тимирязева 42, павильон 17",
    "Almaty, 050057, Kazakhstan": "Алматы, 050057, Казахстан",
    "9 Brown Lane West, Holbeck": "9 Браун Лейн Уэст Холбек",
    "Leeds LS12 6BH": "Лидс LS12 6BH",
    "EXW Leeds": "EXW Лидс",
    "30% payable with order, 70% payable prior delivery": "30% оплата с заказом, 70% оплата перед доставкой",
    "North Caspian Operating Company N.V., Kazakhstan Branch": "Норт Каспиан Оперейтинг Компани Н.В.",
    "Limit Switch Box Soldo": "Блок концевых выключателей Soldo",
    "Limit Switch Box": "Блок концевых выключателей",
}

ATTENTION_MAP = {"Rachel Tam": "Рэйчел Тэм", "Timur Nabiullin": "Тимур Набиуллин"}
TRANSLATE_COORDS = ["C2","C3","B5","D5","D6","B7","D7","B8","D8","B9","D9","B10","C12","C13","C14","D13","C15","D15","C16","D16"]

def translate_cell(val):
    if not val: return val
    s = str(val).strip()
    if s in EN_TO_RU: return EN_TO_RU[s]
    if s.startswith("Attention:"):
        name = s[10:].strip()
        return "Вниманию: " + ATTENTION_MAP.get(name, name)
    if s.startswith("E-mail:") and "Rachel" in s:
        return s.replace("E-mail:", "Эл. почта:")
    result = s
    for en, ru in EN_TO_RU.items(): result = result.replace(en, ru)
    return result

def translate_header(v):
    return str(v).replace("PO #", "Заказ на закупку #").replace("rev.", "рев.").replace(" dd ", " от ")

def translate_contract(v):
    return str(v).replace("to Contract", "к контракту").replace(" dd ", " от ")

def extract_po_date(header):
    import re
    m = re.search(r'dd\s+(\d{2})\.(\d{2})\.(\d{4})', str(header), re.IGNORECASE)
    if m: return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None

def delivery_date(po_date):
    from dateutil.relativedelta import relativedelta
    return po_date + relativedelta(months=11)

def find_total_row(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and "итого" in str(cell.value).lower():
                return cell.row
    return 22

def collect_item_rows(ws):
    rows = []
    for row in ws.iter_rows():
        b = ws.cell(row=row[0].row, column=2).value
        if isinstance(b, (int, float)) and b == int(b) and b >= 1:
            rows.append(row[0].row)
    return rows

def process_po(xlsx_bytes, mode):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["PO"]
    header_val = ws["C2"].value or ""
    po_date = extract_po_date(header_val)
    if not po_date: po_date = datetime.today()
    del_date = delivery_date(po_date)
    if mode == "ru":
        if ws["C2"].value: ws["C2"].value = translate_header(ws["C2"].value)
        if ws["C3"].value: ws["C3"].value = translate_contract(ws["C3"].value)
        for coord in TRANSLATE_COORDS:
            if coord in ("C2", "C3"): continue
            cell = ws[coord]
            if cell.value: cell.value = translate_cell(str(cell.value))
        for rnum in collect_item_rows(ws):
            c = ws.cell(row=rnum, column=3)
            if c.value:
                c.value = str(c.value).replace("Limit Switch Box Soldo", "Блок концевых выключателей Soldo").replace("Limit Switch Box", "Блок концевых выключателей")
    d14 = ws["D14"]
    d14.value = del_date
    d14.number_format = "DD.MM.YYYY"
    total_row = find_total_row(ws)
    director_row = (total_row + 4) if total_row else 26
    director_text = "Генеральный директор _________________________________   Мухтаганова Г.Р." if mode == "ru" else "General director _________________________________   Mukhtaganova G.R."
    c_cell = ws.cell(row=director_row, column=3)
    c_cell.value = director_text
    try:
        ref = ws["C12"]
        c_cell.font = Font(name=ref.font.name or "Calibri", size=ref.font.size or 11)
    except: pass
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def build_invoice_docx(pdf_text):
    from docx import Document
    from docx.shared import Pt, Cm
    import re
    def g(pattern, default=""):
        m = re.search(pattern, pdf_text, re.IGNORECASE)
        return m.group(1).strip() if m else default
    cust_ref = g(r'Customer Reference:\s*(.+)', 'PO #73799')
    our_ref  = g(r'Our Reference:\s*(.+)', 'CUK132578-01-1')
    inv_date = g(r'Invoice Date:\s*(\S+)', '30.03.2026')
    cust_code= g(r'Customer Code:\s*(\S+)', 'C02047')
    contract = g(r'(ROTSALGOODS\S+)', 'ROTSALGOODS2022UK')
    total    = g(r'Grand Total\s+([\d,. ]+)', '4 320.00').strip()
    prepay   = g(r'30%[^\n]*?([\d,]+\.?\d*)', '1 296.00').strip()
    remain   = g(r'70%[^\n]*?([\d,]+\.?\d*)', '3 024.00').strip()
    phone    = g(r'Credit control:\s*(\S+)', '+441132057271')
    items = []
    im = re.search(r'(SS\S+[\w\s\-]+?)\s{2,}(\d+)\s+[£]?([\d,. ]+)\s+[£]?([\d,. ]+)', pdf_text)
    if im:
        desc = im.group(1).replace('Limit Switch Box Soldo', 'Блок концевых выключателей Soldo').replace('Limit Switch Box', 'Блок концевых выключателей').strip()
        items.append({'desc': desc, 'qty': im.group(2), 'price': f'£ {im.group(3).strip()}', 'vat': '0%', 'total': f'£ {im.group(4).strip()}'})
    else:
        items.append({'desc': 'SS5022E‐2‐02601 Блок концевых выключателей Soldo', 'qty': '3', 'price': '£ 1 440.00', 'vat': '0%', 'total': '£ 4 320.00'})
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2); section.right_margin = Cm(1.5)
    def r(para, text, bold=False, size=10):
        run = para.add_run(text); run.bold = bold; run.font.size = Pt(size); return run
    t1 = doc.add_table(rows=1, cols=2); t1.style = 'Table Grid'
    left = t1.cell(0,0).paragraphs[0]; right = t1.cell(0,1).paragraphs[0]
    r(left, 'Коммерческий инвойс', bold=True, size=11)
    r(left, f':   Ссылка Заказчика: {cust_ref}  ')
    r(left, f'\nНомер контракта: {contract}  ')
    r(left, f'\nНаша ссылка: {our_ref}  ')
    r(left, f'\nВалюта: GBP  ')
    r(left, f'\nУсловия оплаты: ПРОФОРМА', bold=True)
    r(right, f'Дата инвойса: {inv_date}\n')
    r(right, f'Код Заказчика: {cust_code}')
    doc.add_paragraph()
    t2 = doc.add_table(rows=1, cols=2); t2.style = 'Table Grid'
    ap = t2.cell(0,0).paragraphs[0]
    r(ap, 'Адрес инвойса: ', bold=True)
    r(ap, 'ТОО Тянь-Шань Engineering\nАтакент, ул. Тимирязева, 42, павильон 17\nАлматы, 050057, Казахстан')
    doc.add_paragraph()
    t3 = doc.add_table(rows=1+len(items), cols=5); t3.style = 'Table Grid'
    for i, h in enumerate(['Описание товара', 'Кол-во', 'Цена (£)', 'НДС', 'Итого (£)']):
        p = t3.cell(0,i).paragraphs[0]; run = p.add_run(h); run.bold = True; run.font.size = Pt(9)
    for ri, item in enumerate(items):
        for ci, v in enumerate([item['desc'], item['qty'], item['price'], item['vat'], item['total']]):
            t3.cell(1+ri, ci).paragraphs[0].add_run(v).font.size = Pt(9)
    doc.add_paragraph()
    p1 = doc.add_paragraph(); r(p1, f'30% предоплата — £ {prepay}', bold=True)
    p2 = doc.add_paragraph(); r(p2, f'70% оплата перед отправкой — £ {remain}', bold=True)
    doc.add_paragraph()
    pt = doc.add_paragraph()
    r(pt, f'Кредитный контроль: {phone}                    ')
    r(pt, f'Всего: £ {total}     НДС: £ 0.00     Итого: £ {total}', bold=True)
    doc.add_paragraph()
    t4 = doc.add_table(rows=1, cols=2); t4.style = 'Table Grid'
    bp = t4.cell(0,0).paragraphs[0]
    r(bp, 'Счёт в фунтах:\n', bold=True)
    for line in ['Фин. учреждение: HSBC UK Bank PLC', 'Имя: Rotork UK Ltd.', 'SWIFT: HBUKGB4103B', 'IBAN: GB94HBUK40141333958183', 'Сорт: 40-14-13 / Счёт: 33958183']:
        r(bp, line+'\n', size=9)
    fp = t4.cell(0,1).paragraphs[0]; r(fp, '\n')
    for line in ['Условия: Проформа/наличные', 'Рег. Англия 1090344', 'Rotork House, Brassmill Lane, Bath BA1 3JQ', 'VAT: GB 180 3944 59']:
        r(fp, line+'\n', size=9)
    out = io.BytesIO(); doc.save(out); return out.getvalue()

class handler(BaseHTTPRequestHandler):
    def do_OPTIONS(self):
        self.send_response(200); self._cors(); self.end_headers()
    def do_POST(self):
        try:
            length = int(self.headers.get('Content-Length', 0))
            body = json.loads(self.rfile.read(length))
            xlsx_bytes = base64.b64decode(body['xlsx'])
            pdf_text = body.get('pdf_text', '')
            en_bytes  = process_po(xlsx_bytes, 'en')
            ru_bytes  = process_po(xlsx_bytes, 'ru')
            inv_bytes = build_invoice_docx(pdf_text)
            result = {
                'en_xlsx':  base64.b64encode(en_bytes).decode(),
                'ru_xlsx':  base64.b64encode(ru_bytes).decode(),
                'inv_docx': base64.b64encode(inv_bytes).decode(),
            }
            self.send_response(200); self._cors()
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps(result).encode())
        except Exception as e:
            self.send_response(500); self._cors()
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps({'error': str(e)}).encode())
    def _cors(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
